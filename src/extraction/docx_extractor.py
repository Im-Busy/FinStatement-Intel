"""DOCX extraction — extract financial statements from Word documents.

Uses python-docx to parse tables and paragraphs, classifying content
into Income Statement, Balance Sheet, and Cash Flow Statement line items.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

from src.extraction.line_item_extractor import build_line_item, clean_label
from src.extraction.unit_detector import detect_unit

logger = logging.getLogger(__name__)

_STATEMENT_KEYWORDS: dict[str, list[str]] = {
    "IS": [
        r"revenue",
        r"sales",
        r"cost of",
        r"gross profit",
        r"operating (?:income|profit)",
        r"net income",
        r"earnings per share",
        r"eps",
        r"ebitda",
        r"income tax",
        r"r&d",
        r"research and development",
        r"selling,? general",
        r"interest (?:income|expense)",
    ],
    "BS": [
        r"total assets",
        r"current assets",
        r"current liabilities",
        r"total liabilities",
        r"shareholders.? equity",
        r"cash and (?:cash )?equivalents",
        r"accounts? receivable",
        r"inventor(?:y|ies)",
        r"property.*equipment",
        r"accounts? payable",
        r"long-term debt",
        r"goodwill",
        r"retained earnings",
        r"accumulated (?:other )?comprehensive",
    ],
    "CFS": [
        r"operating activities",
        r"investing activities",
        r"financing activities",
        r"depreciation and amortization",
        r"amortization",
        r"stock.based compensation",
        r"capital expenditures?",
        r"capex",
        r"free cash flow",
        r"proceeds from",
        r"repayment of",
        r"dividends? paid",
        r"share repurchase",
        r"net change in cash",
    ],
}

_AMOUNT_RE = re.compile(r"^\s*\$?\s*-?\(?[\d,.]+\)?\s*$")
_NUMBER_RE = re.compile(r"[\d,.]+")


def _classify_row(label: str) -> str:
    """Classify a row label into IS, BS, or CFS based on keyword matching.

    Args:
        label: The text label of a row.

    Returns:
        'IS', 'BS', or 'CFS' statement type key.
    """
    lower = label.lower().strip()
    for stmt, patterns in _STATEMENT_KEYWORDS.items():
        for pattern in patterns:
            if re.search(pattern, lower):
                return stmt
    return "BS"


def _extract_table_data(table: Any) -> list[dict[str, Any]]:
    """Extract line items from a python-docx Table object.

    Args:
        table: A python-docx Table instance.

    Returns:
        List of standardized line item dicts.
    """
    line_items: list[dict[str, Any]] = []
    seen_labels: set[str] = set()

    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        if not cells or len(cells) < 2:
            continue

        label = clean_label(cells[0])
        if not label or len(label) < 3:
            continue

        label_lower = label.lower()
        if label_lower in seen_labels:
            continue

        value = 0.0
        for cell_text in cells[1:]:
            cleaned = cell_text.replace("$", "").replace(",", "").strip()
            if cleaned.startswith("(") and cleaned.endswith(")"):
                cleaned = f"-{cleaned[1:-1]}"
            try:
                value = float(cleaned)
                break
            except (ValueError, TypeError):
                continue

        if label_lower in (
            "total assets",
            "total liabilities",
            "total shareholders' equity",
            "total equity",
            "net income",
            "revenue",
            "total revenue",
            "net revenue",
        ):
            seen_labels.add(label_lower)
            line_items.append(build_line_item(label, value))
        elif len(label) > 3:
            seen_labels.add(label_lower)
            line_items.append(build_line_item(label, value))

    return line_items


def _extract_paragraph_text(doc: Any) -> str:
    """Extract all paragraph text from a docx document.

    Args:
        doc: A python-docx Document instance.

    Returns:
        Concatenated paragraph text.
    """
    texts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            texts.append(text)
    return "\n".join(texts)


def extract_from_docx(
    docx_path: str | Path,
    ticker: str = "",
) -> dict[str, Any]:
    """Extract financial statement data from a Word (.docx) document.

    Parses tables and paragraphs, classifies line items by keyword matching,
    and returns standard extraction JSON.

    Args:
        docx_path: Path to the .docx file.
        ticker: Optional company ticker for metadata.

    Returns:
        Standard extraction result dict.
    """
    try:
        from docx import Document
    except ImportError:
        return {
            "status": "ERROR",
            "message": "python-docx not installed. Install with: uv add python-docx",
            "company": {"name": ticker.upper(), "ticker": ticker.upper(), "fiscal_year_end": ""},
            "periods": [],
            "warnings": [],
        }

    path = Path(docx_path)
    if not path.exists():
        return {
            "status": "ERROR",
            "message": f"File not found: {docx_path}",
            "company": {"name": ticker.upper(), "ticker": ticker.upper(), "fiscal_year_end": ""},
            "periods": [],
            "warnings": [],
        }

    try:
        doc = Document(str(path))
    except Exception as e:
        return {
            "status": "ERROR",
            "message": f"Failed to open document: {e}",
            "company": {"name": ticker.upper(), "ticker": ticker.upper(), "fiscal_year_end": ""},
            "periods": [],
            "warnings": [],
        }

    all_line_items: dict[str, list[dict[str, Any]]] = {"IS": [], "BS": [], "CFS": []}

    for table in doc.tables:
        items = _extract_table_data(table)
        for item in items:
            stmt = _classify_row(item["label"])
            all_line_items[stmt].append(item)

    if sum(len(v) for v in all_line_items.values()) == 0:
        return {
            "status": "ERROR",
            "message": f"No financial data tables found in {path.name}",
            "company": {"name": ticker.upper(), "ticker": ticker.upper(), "fiscal_year_end": ""},
            "periods": [],
            "warnings": [],
        }

    full_text = _extract_paragraph_text(doc)
    unit_name, currency = detect_unit([full_text])

    period = {
        "type": "annual",
        "fiscal_year": 0,
        "end_date": "",
        "source": f"docx:{path.name}",
        "statements": {
            "IS": {"unit": unit_name, "line_items": all_line_items["IS"]},
            "BS": {"unit": unit_name, "line_items": all_line_items["BS"]},
            "CFS": {"unit": unit_name, "line_items": all_line_items["CFS"]},
        },
    }

    return {
        "status": "SUCCESS",
        "company": {
            "name": ticker.upper() if ticker else "UNKNOWN",
            "ticker": ticker.upper() if ticker else "",
            "fiscal_year_end": "",
            "currency": currency,
        },
        "periods": [period],
        "warnings": [],
        "metadata": {
            "extraction_method": "docx",
            "extraction_timestamp": "",
            "periods_requested": 1,
            "periods_extracted": 1,
        },
    }
